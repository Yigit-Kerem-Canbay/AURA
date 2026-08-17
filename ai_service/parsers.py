import fitz  # PyMuPDF
import pdfplumber
import docx
import openpyxl
import pytesseract
from PIL import Image
import io

class DocumentParser:
    def parse(self, file_path: str):
        raise NotImplementedError

class PDFParser(DocumentParser):
    def parse(self, file_path: str):
        pages = []
        try:
            # First try PyMuPDF
            doc = fitz.open(file_path)
            for page_num in range(len(doc)):
                page = doc.load_page(page_num)
                text = page.get_text()
                
                # If text is too short, might be a scanned PDF, fallback to OCR
                if len(text.strip()) < 50:
                    text = self._ocr_page(page)
                    
                pages.append({"page_number": page_num + 1, "text": text})
            doc.close()
        except Exception as e:
            print(f"Error parsing PDF with PyMuPDF: {e}")
        return pages

    def _ocr_page(self, page):
        pix = page.get_pixmap()
        img = Image.open(io.BytesIO(pix.tobytes()))
        # Ensure tesseract is installed and turkish language is available
        text = pytesseract.image_to_string(img, lang='tur+eng')
        return text

class DocxParser(DocumentParser):
    def parse(self, file_path: str):
        pages = []
        try:
            doc = docx.Document(file_path)
            full_text = []
            for para in doc.paragraphs:
                full_text.append(para.text)
            
            # DOCX doesn't have a strict page concept, treat as one page or chunk by paragraphs
            pages.append({"page_number": 1, "text": '\n'.join(full_text)})
        except Exception as e:
            print(f"Error parsing DOCX: {e}")
        return pages

class XlsxParser(DocumentParser):
    def parse(self, file_path: str):
        pages = []
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            for sheet_idx, sheet in enumerate(wb.worksheets):
                text = []
                for row in sheet.iter_rows(values_only=True):
                    row_text = ' '.join([str(cell) for cell in row if cell is not None])
                    if row_text:
                        text.append(row_text)
                pages.append({"page_number": sheet_idx + 1, "text": '\n'.join(text)})
        except Exception as e:
            print(f"Error parsing XLSX: {e}")
        return pages

def get_parser(file_type: str) -> DocumentParser:
    file_type = file_type.lower().strip('.')
    if file_type == 'pdf':
        return PDFParser()
    elif file_type == 'docx':
        return DocxParser()
    elif file_type == 'xlsx':
        return XlsxParser()
    else:
        raise ValueError(f"Unsupported file type: {file_type}")
